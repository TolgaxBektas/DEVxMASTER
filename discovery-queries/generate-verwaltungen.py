#!/usr/bin/env python3
import argparse
import json
import random
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = Path(
    "/home/ubuntu/attachments/c3822988-c901-40e6-918f-eda360c86c78/"
    "Judith_Winkler_Gemeinden_PLZ_9_bis_0.xlsx"
)
LOCALITIES_PATH = ROOT / "modules/ingestion/src/data/localities.de.json"
AREAS_PATH = ROOT / "modules/ingestion/src/data/areas.de.json"
OUTPUT_PATH = ROOT / "modules/ingestion/src/data/verwaltungen.de.json"
REPORT_PATH = Path("/home/ubuntu/pruefung/recherche-verwaltungen-generator-report.txt")

EXCLUDED_PREFIXES = (
    "landratsamt",
    "kreisverwaltung",
    "kreisausschuss",
    "landkreis",
    "kreis ",
    "regierung",
    "regierungspräsidium",
    "regierungsbezirk",
    "landesdirektion",
    "landesregierung",
    "staatskanzlei",
    "regionalverband",
)
LOCAL_AUTHORITY_PREFIXES = (
    "stadt",
    "gemeinde",
    "amt",
    "markt",
    "samtgemeinde",
    "verbandsgemeinde",
    "verwaltungsgemeinschaft",
)
SCOPE_PREFIXES = (
    "Gemeindevorstand der Gemeinde ",
    "Magistrat der Stadt ",
    "Magistrat der Gemeinde ",
    "Gemeindeverwaltung ",
    "Stadtverwaltung ",
    "Große Kreisstadt ",
    "Kreisfreie Stadt ",
    "Landeshauptstadt ",
    "Gemeinde ",
    "Stadt ",
    "Markt ",
    "Flecken ",
    "EG ",
)


def clean_name(value: Any) -> str:
    return "" if value is None else " ".join(str(value).split())


def normalize_postal_code(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if text.isdigit():
        return text.zfill(5)
    return text


def exclusion_reason(name: str) -> str | None:
    if not name:
        return "leerer Name"
    folded_name = name.casefold()
    for prefix in sorted(EXCLUDED_PREFIXES, key=len, reverse=True):
        if folded_name.startswith(prefix):
            return f"Name-Präfix {prefix.strip()}"
    if not folded_name.startswith(LOCAL_AUTHORITY_PREFIXES):
        last_word = folded_name.split()[-1].rstrip(".,;:")
        if last_word.endswith("kreis"):
            return "Name endet auf Kreis"
    return None


def normalize_scope(name: str) -> str:
    value = clean_name(name)
    folded = value.casefold()
    if folded.startswith("vereinbarte verwaltungsgemeinschaft "):
        value = "Verwaltungsgemeinschaft " + value[len("Vereinbarte Verwaltungsgemeinschaft "):]
    elif folded.startswith("verwaltungsgemeinschaft der stadt "):
        value = "Verwaltungsgemeinschaft " + value[len("Verwaltungsgemeinschaft der Stadt "):]
    elif folded.startswith("verbandsgemeindeverwaltung "):
        value = "Verbandsgemeinde " + value[len("Verbandsgemeindeverwaltung "):]
    elif folded.startswith("vg "):
        value = "Verwaltungsgemeinschaft " + value[3:]
    else:
        for prefix in SCOPE_PREFIXES:
            if value.casefold().startswith(prefix.casefold()):
                value = value[len(prefix):]
                break
    value = re.sub(r",\s*(?:Stadt|Gemeinde)\s*$", "", value, flags=re.IGNORECASE)
    return clean_name(value)


def locality_aggs(localities: list[dict]) -> tuple[dict, dict]:
    exact: dict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    postal: dict[str, Counter[str]] = defaultdict(Counter)
    for locality in localities:
        postal_code = normalize_postal_code(locality.get("postalCode"))
        name = clean_name(locality.get("name"))
        municipality_key = clean_name(locality.get("municipalityKey"))
        if len(municipality_key) < 5:
            continue
        ags = municipality_key[:5]
        postal[postal_code][ags] += 1
        exact[(postal_code, name)][ags] += 1
    return exact, postal


def select_ags(counts: Counter[str]) -> str | None:
    if not counts:
        return None
    return min(counts, key=lambda ags: (-counts[ags], ags))


def normalize_comparison_name(value: str) -> str:
    normalized = clean_name(value).casefold()
    normalized = re.sub(r"\([^)]*\)", " ", normalized)
    for pattern, replacement in (
        (r"\ba\s*\.\s*d\s*\.\s*", "an der "),
        (r"\bi\s*\.\s*d\s*\.\s*", "in der "),
        (r"\bv\s*\.\s*d\s*\.\s*", "vor der "),
        (r"\ba\s*\.\s*", "am "),
        (r"\bi\s*\.\s*", "im "),
        (r"\bb\s*\.\s*", "bei "),
        (r"\bst\s*\.\s*", "sankt "),
    ):
        normalized = re.sub(pattern, replacement, normalized)
    return " ".join(normalized.replace(".", " ").split())


def build_mapping_indexes(
    localities: list[dict],
    area_by_ags: dict[str, dict],
) -> dict[str, Any]:
    exact, postal = locality_aggs(localities)
    names_by_state: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    normalized_by_state: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    postal_by_ags: dict[str, list[str]] = defaultdict(list)
    for locality in localities:
        ags = clean_name(locality.get("municipalityKey"))[:5]
        area = area_by_ags.get(ags)
        if not area:
            continue
        state_name = clean_name(area.get("stateName"))
        name = clean_name(locality.get("name"))
        if not state_name or not name:
            continue
        state_key = state_name.casefold()
        names_by_state[state_key][name.casefold()].add(ags)
        normalized_by_state[state_key][normalize_comparison_name(name)].add(ags)
        postal_by_ags[ags].append(normalize_postal_code(locality.get("postalCode")))
    return {
        "exact": exact,
        "postal": postal,
        "names_by_state": names_by_state,
        "normalized_by_state": normalized_by_state,
        "postal_by_ags": postal_by_ags,
    }


def common_postal_prefix(left: str, right: str) -> int:
    length = 0
    for left_digit, right_digit in zip(left, right):
        if left_digit != right_digit:
            break
        length += 1
    return length


def map_district(
    postal_code: str,
    ort: str,
    scope: str,
    state_name: str,
    indexes: dict[str, Any],
) -> tuple[str | None, str | None]:
    exact_ags = select_ags(indexes["exact"].get((postal_code, ort), Counter()))
    if exact_ags:
        return exact_ags, "a"

    state_key = clean_name(state_name).casefold()
    names_by_state = indexes["names_by_state"].get(state_key, {})
    name_matches = names_by_state.get(clean_name(ort).casefold(), set())
    if len(name_matches) == 1:
        return next(iter(name_matches)), "b"

    normalized_by_state = indexes["normalized_by_state"].get(state_key, {})
    ambiguous_matches = set(name_matches)
    for value in (ort, scope):
        matches = normalized_by_state.get(normalize_comparison_name(value), set())
        if len(matches) == 1:
            return next(iter(matches)), "c"
        ambiguous_matches.update(matches)

    if len(ambiguous_matches) > 1 and postal_code:
        prefix_scores = {
            ags: max(
                (
                    common_postal_prefix(postal_code, locality_postal)
                    for locality_postal in indexes["postal_by_ags"].get(ags, [])
                ),
                default=0,
            )
            for ags in ambiguous_matches
        }
        best_score = max(prefix_scores.values(), default=0)
        winners = [ags for ags, score in prefix_scores.items() if score == best_score]
        if best_score >= 3 and len(winners) == 1:
            return winners[0], "d"

    postal_ags = select_ags(indexes["postal"].get(postal_code, Counter()))
    return (postal_ags, "e") if postal_ags else (None, None)


def generate(input_path: Path) -> tuple[list[dict], list[str]]:
    localities = json.loads(LOCALITIES_PATH.read_text(encoding="utf-8"))
    areas = json.loads(AREAS_PATH.read_text(encoding="utf-8"))
    area_by_ags = {
        area["ags"]: area
        for area in areas
        if area["level"] == "district"
    }
    mapping_indexes = build_mapping_indexes(localities, area_by_ags)
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook["Tabelle1"]
    headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    required = ("Bundesland", "Kreis", "Verwaltung", "Postleitzahl", "Ort")
    positions = {name: headers.index(name) for name in required}
    candidates: list[dict] = []
    exclusions: Counter[str] = Counter()
    excluded_names: dict[str, Counter[str]] = defaultdict(Counter)
    source_name_counts: Counter[str] = Counter()
    for row in worksheet.iter_rows(min_row=2, max_col=6, values_only=True):
        state_name = clean_name(row[positions["Bundesland"]])
        typ = clean_name(row[positions["Kreis"]])
        name = clean_name(row[positions["Verwaltung"]])
        postal_code = normalize_postal_code(row[positions["Postleitzahl"]])
        ort = clean_name(row[positions["Ort"]])
        if not any((state_name, typ, name, postal_code, ort)):
            continue
        if name:
            source_name_counts[name] += 1
        reason = exclusion_reason(name)
        if reason:
            exclusions[reason] += 1
            excluded_names[reason][name or "<leer>"] += 1
            continue
        candidates.append({
            "name": name,
            "scope": normalize_scope(name),
            "typ": typ,
            "postalCode": postal_code,
            "ort": ort,
            "stateName": state_name,
        })
    workbook.close()

    candidates.sort(key=lambda item: (
        item["name"].casefold(),
        item["stateName"].casefold(),
        item["typ"].casefold(),
        item["postalCode"],
        item["ort"].casefold(),
    ))
    unique: dict[str, dict] = {}
    duplicate_names: Counter[str] = Counter()
    duplicate_values: dict[str, list[dict]] = defaultdict(list)
    for candidate in candidates:
        name = candidate["name"]
        if name in unique:
            duplicate_names[name] += 1
            duplicate_values[name].append(candidate)
            continue
        unique[name] = candidate

    mapped: list[dict] = []
    unmapped: list[str] = []
    mapped_by_step: Counter[str] = Counter()
    for candidate in unique.values():
        district_ags, mapping_step = map_district(
            candidate["postalCode"],
            candidate["ort"],
            candidate["scope"],
            candidate["stateName"],
            mapping_indexes,
        )
        if district_ags is None or district_ags not in area_by_ags:
            unmapped.append(
                f"{candidate['name']} | {candidate['postalCode'] or '<leer>'} | "
                f"{candidate['ort'] or '<leer>'}"
            )
            continue
        mapped_by_step[mapping_step or ""] += 1
        mapped.append({
            **candidate,
            "districtAgs": district_ags,
            "mappingStep": mapping_step,
        })

    mapped.sort(key=lambda item: (
        item["districtAgs"],
        item["postalCode"],
        item["scope"].casefold(),
        item["name"].casefold(),
    ))
    lines = [
        f"Quellzeilen mit Daten in den Nicht-Personenspalten: {len(candidates) + sum(exclusions.values())}",
        f"Ausgeschlossen gesamt: {sum(exclusions.values())}",
        "Doppelte Verwaltung-Namen in der Quelle: "
        f"{sum(count - 1 for count in source_name_counts.values() if count > 1)} Zusatzzeilen / "
        f"{sum(1 for count in source_name_counts.values() if count > 1)} Namen",
    ]
    for reason in sorted(exclusions):
        lines.append(f"  {reason}: {exclusions[reason]}")
        lines.extend(
            f"    {name}" + (f" (×{count})" if count > 1 else "")
            for name, count in sorted(excluded_names[reason].items(), key=lambda item: item[0].casefold())
        )
    lines.extend([
        f"Doppelte Namen nach Ausschlüssen zusammengeführt: {sum(duplicate_names.values())} Zusatzzeilen / {len(duplicate_names)} Namen",
        f"Eindeutige Verwaltungen nach Ausschlüssen: {len(unique)}",
        f"Unmapped: {len(unmapped)}",
        "Zuordnungen nach Mapping-Schritt:",
        *(
            f"  Schritt {step}: {mapped_by_step[step]}"
            for step in "abcde"
        ),
    ])
    if duplicate_names:
        lines.append("Zusammengeführte Doppelnamen:")
        lines.extend(
            f"  {name} (zusätzliche Zeilen: {count})"
            for name, count in sorted(duplicate_names.items(), key=lambda item: item[0].casefold())
        )
    lines.append("Nicht zuordenbare Verwaltungen (Name | PLZ | Ort):")
    lines.extend(f"  {row}" for row in sorted(unmapped, key=str.casefold))
    rng = random.Random(20261012)
    cascade_mapped = [row for row in mapped if row["mappingStep"] in {"b", "c", "d", "e"}]
    sample_cascade = rng.sample(cascade_mapped, min(10, len(cascade_mapped)))
    remaining = [row for row in mapped if row not in sample_cascade]
    sample = [*sample_cascade, *rng.sample(remaining, min(40 - len(sample_cascade), len(remaining)))]
    rng.shuffle(sample)
    lines.append(
        "Zufallsstichprobe (Seed 20261012), mindestens 10 Zuordnungen aus b–e; "
        "Name → scope → districtAgs / Kreis / Schritt:"
    )
    lines.extend(
        f"  {row['name']} → {row['scope']} → {row['districtAgs']} / "
        f"{area_by_ags[row['districtAgs']]['name']} / {row['mappingStep']}"
        for row in sample
    )
    return mapped, lines


def main() -> None:
    parser = argparse.ArgumentParser(description="Generiert bereinigte Verwaltungssuchbereiche.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    parser.add_argument("--report", type=Path, default=REPORT_PATH)
    args = parser.parse_args()
    mapped, report_lines = generate(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(mapped, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.report.write_text("\n".join([f"Verwaltungen übernommen: {len(mapped)}", *report_lines]) + "\n", encoding="utf-8")
    print(f"verwaltungen_kept={len(mapped)}")
    print(f"report={args.report}")
    print(f"output={args.output}")


if __name__ == "__main__":
    main()
